"""Tests for `scripts/export_pbp_csv.py`.

Scoped to the script's own pure/writer logic (column reordering, the utf-8-sig CSV writer via
`flag_football_ep.owner_csv`, the xlsx writer) -- `build_korpus_frame`/`build_wm_frame`
themselves need a full `Config` and the real ingest pipeline and are exercised end-to-end
elsewhere (`tests/test_ingest_ifaf.py`, `tests/test_pipeline_ingest.py`), not re-tested here.
Every fixture is synthetic, no real player names.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import polars as pl

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "scripts"))

import export_pbp_csv as exp  # noqa: E402

from flag_football_ep.owner_csv import read_owner_csv  # noqa: E402


# --- _select_front_first -----------------------------------------------------------------


def test_select_front_first_reorders_present_columns_only() -> None:
    df = pl.DataFrame({"z": [1], "yards_gained": [5], "ep": [0.5], "a": [9]})

    out = exp._select_front_first(df, exp._FRONT_COLUMNS)

    # yards_gained then ep (both in _FRONT_COLUMNS, present), then every remaining column
    # (z, a) in their original relative order.
    assert out.columns == ["yards_gained", "ep", "z", "a"]


def test_select_front_first_skips_absent_front_columns() -> None:
    df = pl.DataFrame({"only_col": [1]})

    out = exp._select_front_first(df, exp._FRONT_COLUMNS)

    assert out.columns == ["only_col"]


def test_select_front_first_preserves_row_values() -> None:
    df = pl.DataFrame({"b": [1, 2], "yards_gained": [10, 20]})

    out = exp._select_front_first(df, exp._FRONT_COLUMNS)

    assert out["yards_gained"].to_list() == [10, 20]
    assert out["b"].to_list() == [1, 2]


# --- write_xlsx ----------------------------------------------------------------------------


def test_write_xlsx_single_sheet_header_and_rows(tmp_path: Path) -> None:
    df = pl.DataFrame({"game_id": ["g1", "g2"], "note": ["Nühse", "plain"]})
    path = tmp_path / "out.xlsx"

    exp.write_xlsx(df, path)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["pbp"]
    ws = wb["pbp"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("game_id", "note")
    assert rows[1] == ("g1", "Nühse")
    assert rows[2] == ("g2", "plain")


def test_write_xlsx_creates_parent_directories(tmp_path: Path) -> None:
    df = pl.DataFrame({"a": [1]})
    path = tmp_path / "nested" / "dir" / "out.xlsx"

    exp.write_xlsx(df, path)

    assert path.exists()


# --- CSV writers use write_owner_csv (utf-8-sig BOM) ----------------------------------------


def test_build_korpus_frame_output_written_via_owner_csv_has_bom(tmp_path: Path) -> None:
    """The script's own writer path (`write_owner_csv`) is exercised directly here rather
    than through `main`, which needs a full pipeline `Config` -- confirms the exact call
    `main` makes for `korpus_alle_quellen_pbp.csv` produces a BOM'd, round-trippable file."""
    from flag_football_ep.owner_csv import write_owner_csv

    df = pl.DataFrame({"game_id": ["g1"], "note": ["Nühse"]})
    path = tmp_path / "korpus_alle_quellen_pbp.csv"

    write_owner_csv(df, path)

    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows, notices = read_owner_csv(path)
    assert rows == [{"game_id": "g1", "note": "Nühse"}]
    assert notices == []
